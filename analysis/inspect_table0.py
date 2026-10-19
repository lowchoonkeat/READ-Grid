import docx
doc = docx.Document('paper_package/template/iem_template_original_backup.docx')
t0 = doc.tables[0]
c1 = t0.rows[0].cells[1]
print('c1 paragraphs count:', len(c1.paragraphs))
for i, p in enumerate(c1.paragraphs):
    print(f'P{i}: text={p.text!r}, pBdr={len(p._p.xpath("./w:pPr/w:pBdr"))}')
nt = c1.tables[0]
print('nested table paragraphs count:', len(nt.rows[0].cells[0].paragraphs))
for i, p in enumerate(nt.rows[0].cells[0].paragraphs):
    print(f'NT P{i}: text={p.text!r}')
    for r in p.runs:
        print(f'   run: text={r.text!r}, bold={r.bold}, italic={r.italic}')
