import docx
import xml.etree.ElementTree as ET

doc = docx.Document('paper_package/template/iem_template_original_backup.docx')

print("=== P0 to P6 ===")
for i in range(7):
    p = doc.paragraphs[i]
    runs_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.font.superscript) for r in p.runs]
    pPr = p._p.pPr.xml if p._p.pPr is not None else ''
    print(f"P{i}: text={p.text!r}")
    print(f"     runs={runs_info}")
    print(f"     pPr={pPr}")

print("\n=== Table 0 ===")
t0 = doc.tables[0]
tblPr = t0._tbl.tblPr
borders = tblPr.xpath('./w:tblBorders')
if borders:
    print("t0 borders:", ET.tostring(borders[0], encoding='utf-8').decode('utf-8'))
for r_i, r in enumerate(t0.rows):
    for c_i, c in enumerate(r.cells):
        print(f"Cell ({r_i},{c_i}) width={c.width.pt}pt")
        print(f"  tcPr={c._tc.tcPr.xml if c._tc.tcPr is not None else ''}")
        for p_idx, p in enumerate(c.paragraphs):
            print(f"    p{p_idx}: align={p.alignment} space_before={p.paragraph_format.space_before} space_after={p.paragraph_format.space_after} line_spacing={p.paragraph_format.line_spacing}")
            print(f"      runs={[(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in r.runs] if hasattr(r, 'runs') else [(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]}")

print("\n=== CRediT Table (Table 11) ===")
t11 = doc.tables[11]
tblPr11 = t11._tbl.tblPr
borders11 = tblPr11.xpath('./w:tblBorders')
if borders11:
    print("t11 borders:", ET.tostring(borders11[0], encoding='utf-8').decode('utf-8'))
for r_i, r in enumerate(t11.rows):
    print(f"Row {r_i}: {[c.text.strip() for c in r.cells]} widths={[c.width.pt for c in r.cells]}")
    for c_i, c in enumerate(r.cells):
        for p in c.paragraphs:
            print(f"  ({r_i},{c_i}) pPr={p._p.pPr.xml if p._p.pPr is not None else ''}")
            for run in p.runs:
                print(f"    run: text={run.text!r}, font={run.font.name}, size={run.font.size.pt if run.font.size else None}, bold={run.bold}")

# Also let's inspect the paragraph right before and after Table 11
print("\n=== Where is Table 11 in body? ===")
body_children = list(doc.element.body)
for idx, ch in enumerate(body_children):
    if ch is t11._tbl:
        print(f"Table 11 is at body child index {idx}")
        # check previous 2 and next 2 children
        for rel in range(-2, 3):
            target_idx = idx + rel
            if 0 <= target_idx < len(body_children):
                target = body_children[target_idx]
                tag = target.tag.split('}')[-1]
                if tag == 'p':
                    p_obj = docx.text.paragraph.Paragraph(target, doc)
                    print(f"  Rel {rel} (idx {target_idx}): [P] {p_obj.text!r}")
                elif tag == 'tbl':
                    print(f"  Rel {rel} (idx {target_idx}): [TBL]")
                elif tag == 'sectPr':
                    print(f"  Rel {rel} (idx {target_idx}): [sectPr]")
