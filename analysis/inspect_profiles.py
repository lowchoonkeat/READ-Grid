import docx, xml.etree.ElementTree as ET
doc = docx.Document('paper_package/template/iem_template_original_backup.docx')

print('Total tables:', len(doc.tables))
t14 = doc.tables[14]
print('=== TABLE 14 (Banner) ===')
print('Rows:', len(t14.rows), 'Cols:', len(t14.columns))
print('tblPr:', t14._tbl.tblPr.xml)
print('cell text:', repr(t14.rows[0].cells[0].text))
print('cell tcPr:', t14.rows[0].cells[0]._tc.tcPr.xml)
for p in t14.rows[0].cells[0].paragraphs:
    print('  p:', p.text, [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.font.color.rgb if r.font.color else None) for r in p.runs])

t15 = doc.tables[15]
print('\n=== TABLE 15 (Profiles) ===')
print('Rows:', len(t15.rows), 'Cols:', len(t15.columns))
print('tblPr:', t15._tbl.tblPr.xml)
for r_i, r in enumerate(t15.rows):
    print(f'Row {r_i}: widths={[c.width.pt for c in r.cells]}')
    cell0 = r.cells[0]
    cell1 = r.cells[1]
    print(f'  cell0 text={cell0.text!r} num_p={len(cell0.paragraphs)}')
    for p in cell0.paragraphs:
        drawings = p._p.xpath('.//w:drawing')
        print(f'    drawings in cell0: {len(drawings)}')
        # check drawing extent / size
        for d in drawings:
            extent = d.xpath('.//wp:extent', namespaces={'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'})
            if extent:
                cx = extent[0].get('cx')
                cy = extent[0].get('cy')
                print(f'      image size: cx={cx} (cx_in={int(cx)/914400:.2f}in, cx_pt={int(cx)/12700:.1f}pt), cy={cy} (cy_in={int(cy)/914400:.2f}in, cy_pt={int(cy)/12700:.1f}pt)')
    print(f'  cell1 text={cell1.text[:60]!r}')
    for p in cell1.paragraphs:
        print(f'    pPr={p._p.pPr.xml if p._p.pPr is not None else None}')
        print(f'    runs={[(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]}')
